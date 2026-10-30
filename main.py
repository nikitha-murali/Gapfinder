import io, os, json, re, logging, time
import requests
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from pydantic import BaseModel

load_dotenv()
BASE = os.path.dirname(os.path.abspath(__file__))
KEY = os.getenv("TINYFISH_API_KEY")
DEMO_ONLY = os.getenv("DEMO_ONLY") == "1"
API = "https://agent.tinyfish.ai/v1/automation/run-sse"
PREVIEW_CACHE_SECONDS = 900
PREVIEW_TIMEOUT_SECONDS = 15
preview_cache = {"expires_at": 0.0, "payload": None}
logger = logging.getLogger(__name__)

BATCH = 4
CHALLENGE_BATCH = 6

SOURCES = {
    "cs": "https://internshala.com/internships/computer-science-internship",
    "web": "https://internshala.com/internships/web-development-internship",
    "data": "https://internshala.com/internships/data-science-internship",
}
CHALLENGE_URL = "https://unstop.com/competitions"

INTERN_GOAL = (
    "Return the first {n} internships on this page as JSON with the key 'internships'. "
    "For each one open its page and return: title, company, stipend, deadline, "
    "apply_url (link to that internship's page), skills (list of required skills), "
    "and summary (2 sentences describing what the intern will actually do)."
)
CHALLENGE_GOAL = (
    "Return up to {n} currently open company-sponsored hackathons, case competitions, "
    "or real-world product, engineering, or business problem challenges as JSON with the key 'challenges'. "
    "Use only information visible on the original listing; do not invent company names, skills, rewards, "
    "or dates. For each one return: title, company (the named organiser, if listed), reward (prize or benefit, "
    "if listed), deadline (if listed), apply_url (the original challenge page), skills (skills explicitly "
    "required or clearly relevant to the stated prompt), and summary (a concise description of the actual "
    "task or problem participants are asked to solve)."
)

MODES = {
    "internships": {"file": "sample.json", "key": "internships", "goal": INTERN_GOAL},
    "challenges": {"file": "sample_challenges.json", "key": "challenges", "goal": CHALLENGE_GOAL},
}

# canonical skill -> other ways people write it
SKILLS = {
    "python": ["python3"],
    "java": [],
    "javascript": ["js", "ecmascript"],
    "typescript": [],
    "c++": ["cpp"],
    "c#": ["csharp"],
    "golang": [],
    "rust": [],
    "kotlin": [],
    "swift": [],
    "php": [],
    "ruby": [],
    "scala": [],
    "matlab": [],
    "r programming": [],
    "html": ["html5", "html 5"],
    "css": ["css3"],
    "react": ["react.js", "reactjs"],
    "angular": ["angularjs"],
    "vue": ["vue.js", "vuejs"],
    "next.js": ["nextjs"],
    "node.js": ["nodejs"],
    "express": ["express.js", "expressjs"],
    "django": [],
    "flask": [],
    "fastapi": [],
    "spring boot": ["springboot"],
    "tailwind css": ["tailwind"],
    "bootstrap": [],
    "sql": [],
    "mysql": [],
    "postgresql": ["postgres"],
    "mongodb": ["mongo db"],
    "sqlite": [],
    "firebase": [],
    "redis": [],
    "git": [],
    "github": [],
    "docker": [],
    "kubernetes": [],
    "linux": [],
    "aws": ["amazon web services"],
    "azure": [],
    "gcp": ["google cloud"],
    "ci/cd": ["cicd"],
    "pandas": [],
    "numpy": [],
    "scikit-learn": ["sklearn"],
    "tensorflow": [],
    "pytorch": [],
    "machine learning": ["ml"],
    "deep learning": [],
    "nlp": ["natural language processing"],
    "computer vision": [],
    "opencv": [],
    "data analysis": ["data analytics"],
    "data visualization": ["data visualisation"],
    "power bi": ["powerbi"],
    "tableau": [],
    "excel": ["ms excel", "microsoft excel"],
    "statistics": [],
    "rest api": ["rest apis", "restful api", "restful apis"],
    "graphql": [],
    "ui/ux design": ["ui & ux design", "ui/ux", "ux design", "ui design"],
    "figma": [],
    "software testing": ["manual testing", "qa testing"],
    "selenium": [],
    "postman": [],
    "android development": ["android"],
    "flutter": [],
    "react native": [],
    "cybersecurity": ["cyber security"],
    "networking": [],
    "blockchain": [],
    "ms-office": ["ms office", "microsoft office"],
    "project management": [],
    "problem solving": [],
    "communication": ["communication skills"],
    "teamwork": [],
    "leadership": [],
    "content writing": [],
    "seo": [],
    "digital marketing": [],
    "social media marketing": [],
    "video editing": [],
    "business development": [],
    "public speaking": [],
}

# having the key skill also covers these
IMPLIES = {
    "react": ["javascript"],
    "react native": ["react", "javascript"],
    "next.js": ["react", "javascript"],
    "node.js": ["javascript"],
    "express": ["node.js", "javascript"],
    "typescript": ["javascript"],
    "angular": ["javascript", "typescript"],
    "vue": ["javascript"],
    "django": ["python"],
    "flask": ["python"],
    "fastapi": ["python"],
    "pandas": ["python", "data analysis"],
    "numpy": ["python"],
    "scikit-learn": ["python", "machine learning"],
    "tensorflow": ["python", "machine learning"],
    "pytorch": ["python", "machine learning"],
    "deep learning": ["machine learning"],
    "mysql": ["sql"],
    "postgresql": ["sql"],
    "sqlite": ["sql"],
    "spring boot": ["java"],
    "github": ["git"],
    "tailwind css": ["css"],
    "bootstrap": ["css", "html"],
    "flutter": ["mobile development"],
}


def _build_aliases():
    out = {}
    for canon, alts in SKILLS.items():
        out[canon] = canon
        for a in alts:
            out[a] = canon
    return out


ALIASES = _build_aliases()


def _pattern(terms):
    alts = "|".join(re.escape(t) for t in sorted(set(terms), key=len, reverse=True))
    return re.compile(r"(?<![a-z0-9+#])(?:" + alts + r")(?![a-z0-9+#])")


PATTERNS = [(canon, _pattern([canon] + alts)) for canon, alts in SKILLS.items()]


def extract_skills(text):
    low = text.lower()
    found = []
    for canon, pat in PATTERNS:
        m = pat.search(low)
        if m:
            found.append((m.start(), canon))
    found.sort()
    return [c for _, c in found]


app = FastAPI()


class Query(BaseModel):
    source: str = "cs"
    demo: bool = True
    mode: str = "internships"
    page: int = 1


def sse(d):
    return f"data: {json.dumps(d)}\n\n"


def stream(q: Query):
    page = max(1, min(q.page, 30))
    is_ch = q.mode == "challenges"
    cfg = MODES["challenges" if is_ch else "internships"]

    if q.demo or DEMO_ONLY:
        items = []
        if page == 1:
            with open(os.path.join(BASE, cfg["file"]), encoding="utf-8") as f:
                items = json.load(f)[cfg["key"]]
        yield sse({"type": "result", "items": items, "demo": True, "page": page})
        return

    if not KEY:
        yield sse({"type": "error", "text": "Server has no TINYFISH_API_KEY set."})
        return

    base = CHALLENGE_URL if is_ch else SOURCES.get(q.source, SOURCES["cs"])
    url = base if (is_ch or page == 1) else f"{base}/page-{page}"
    n = CHALLENGE_BATCH if is_ch else BATCH
    goal = cfg["goal"].format(n=n)

    yield sse({"type": "progress", "text": "Starting the web agent..."})
    try:
        r = requests.post(API, headers={"X-API-Key": KEY},
                          json={"url": url, "goal": goal}, stream=True, timeout=900)
    except Exception as e:
        yield sse({"type": "error", "text": str(e)})
        return
    if r.status_code != 200:
        yield sse({"type": "error", "text": f"TinyFish returned HTTP {r.status_code}"})
        return

    for line in r.iter_lines():
        if not line:
            continue
        text = line.decode("utf-8")
        if not text.startswith("data:"):
            continue
        try:
            ev = json.loads(text[5:].strip())
        except Exception:
            continue
        t = ev.get("type")
        if t == "PROGRESS":
            yield sse({"type": "progress", "text": ev.get("purpose", "")})
        elif t == "COMPLETE":
            if ev.get("status") == "COMPLETED":
                data = ev.get("result")
                if isinstance(data, str):
                    data = json.loads(data)
                items = [i for i in (data or {}).get(cfg["key"], []) if isinstance(i, dict)]
                yield sse({"type": "result", "items": items, "demo": False, "page": page})
            else:
                yield sse({"type": "error", "text": f"Run ended: {ev.get('status')}"})
            return
    yield sse({"type": "error", "text": "The agent stopped without a result."})


@app.post("/api/search")
def search(q: Query):
    return StreamingResponse(stream(q), media_type="text/event-stream")


def sample_listing_preview():
    with open(os.path.join(BASE, "sample.json"), encoding="utf-8") as f:
        return json.load(f)["internships"][:3]


@app.get("/api/listings")
def listings_preview():
    now = time.monotonic()
    cached = preview_cache["payload"]
    if cached is not None and now < preview_cache["expires_at"]:
        return {**cached, "cached": True}

    payload = {"items": sample_listing_preview(), "source": "sample"}
    if KEY and not DEMO_ONLY:
        deadline = now + PREVIEW_TIMEOUT_SECONDS
        try:
            with requests.post(
                API,
                headers={"X-API-Key": KEY},
                json={
                    "url": SOURCES["cs"],
                    "goal": INTERN_GOAL.format(n=3),
                },
                stream=True,
                timeout=(4, 4),
            ) as response:
                response.raise_for_status()
                for line in response.iter_lines():
                    if time.monotonic() >= deadline:
                        logger.warning("TinyFish listing preview timed out; serving sample listings.")
                        break
                    if not line or not line.startswith(b"data:"):
                        continue
                    try:
                        event = json.loads(line[5:].strip())
                    except json.JSONDecodeError:
                        continue
                    if not isinstance(event, dict):
                        continue
                    if event.get("type") != "COMPLETE" or event.get("status") != "COMPLETED":
                        continue
                    result = event.get("result")
                    if isinstance(result, str):
                        try:
                            result = json.loads(result)
                        except json.JSONDecodeError:
                            result = None
                    items = result.get("internships", []) if isinstance(result, dict) else []
                    if items:
                        payload = {
                            "items": [item for item in items if isinstance(item, dict)][:3],
                            "source": "tinyfish",
                        }
                    break
        except requests.RequestException as exc:
            logger.warning("TinyFish listing preview failed; serving sample listings: %s", exc)

    preview_cache["payload"] = payload
    preview_cache["expires_at"] = time.monotonic() + PREVIEW_CACHE_SECONDS
    return {**payload, "cached": False}


@app.post("/api/resume")
async def resume(file: UploadFile | None = File(None), text: str = Form("")):
    raw = text or ""
    if file is not None and file.filename:
        data = await file.read()
        if len(data) > 5_000_000:
            return JSONResponse({"error": "File is too large (max 5 MB)."}, status_code=400)
        name = file.filename.lower()
        try:
            if name.endswith(".pdf"):
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(data))
                raw = "\n".join((p.extract_text() or "") for p in reader.pages)
            elif name.endswith(".docx"):
                import docx
                d = docx.Document(io.BytesIO(data))
                parts = [p.text for p in d.paragraphs]
                for table in d.tables:
                    for row in table.rows:
                        parts.append(" | ".join(c.text for c in row.cells))
                raw = "\n".join(parts)
            elif name.endswith((".txt", ".md")):
                raw = data.decode("utf-8", errors="ignore")
            else:
                return JSONResponse({"error": "Please upload a PDF, DOCX or TXT file."}, status_code=400)
        except Exception:
            return JSONResponse({"error": "Could not read that file. Try pasting the text instead."}, status_code=400)
    raw = raw.strip()
    if not raw:
        return JSONResponse({"error": "No text found. If the PDF is a scan, paste the text instead."}, status_code=400)
    return {"text": raw[:30000], "skills": extract_skills(raw)}


@app.get("/api/config")
def config():
    return {"demo_only": DEMO_ONLY, "aliases": ALIASES, "implies": IMPLIES}


@app.get("/")
def home():
    return FileResponse(os.path.join(BASE, "static", "index.html"))